"""Extraction hardening: text decoding, HTML, zip bombs, YouTube ids, size caps,
legacy Office files."""

import codecs
import gzip
import io
import struct
import time
import zipfile
from pathlib import Path

import pytest

from app import extract, legacy
from app.extract import (
    ExtractError,
    decode_text,
    extract_bytes,
    extract_docx,
    html_to_text,
    youtube_video_id,
)
from app.moodle import MoodleClient, MoodleError, TokenRejected


@pytest.mark.parametrize("blob", [
    codecs.BOM_UTF8 + "Größe".encode(),
    "Größe".encode("utf-16"),  # BOM + LE
    codecs.BOM_UTF16_BE + "Größe".encode("utf-16-be"),
    "Größe".encode("utf-32"),
    "Größe".encode(),
])
def test_decode_text_honours_boms(blob):
    assert decode_text(blob) == "Größe"


def test_decode_text_falls_back_to_latin1_and_drops_nuls():
    assert decode_text("café".encode("latin-1")) == "café"
    assert decode_text(b"a\x00b") == "ab"


def test_text_files_decode_through_extract_bytes():
    assert extract_bytes("notes".encode("utf-16"), "text/plain", "n.txt") == "notes"


def test_html_to_text_drops_markup():
    html = ("<html><head><title>T</title><style>p{color:red}</style></head><body>"
            "<h1>Trees</h1><p>A &amp; B<br>root</p><script>alert(1)</script>"
            "<ul><li>leaf</li><li>node</li></ul>"
            "<table><tr><td>a</td><td>b</td></tr></table></body></html>")
    text = html_to_text(html)
    assert "<" not in text and "alert" not in text and "color" not in text
    assert text == "Trees\n\nA & B\nroot\n\n- leaf\n- node\n\na | b"


@pytest.mark.parametrize("mime, name", [("text/html", "x"), (None, "page.html"),
                                        (None, "page.HTM")])
def test_html_files_are_converted(mime, name):
    assert extract_bytes(b"<p>Hello <b>world</b></p>", mime, name) == "Hello world"


def _zip(members: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in members.items():
            z.writestr(name, data)
    return buf.getvalue()


def test_zip_bomb_docx_is_refused_before_parsing(monkeypatch):
    monkeypatch.setattr(extract, "MAX_UNZIPPED_BYTES", 1000)
    bomb = _zip({"word/document.xml": b"\0" * 5000})
    assert len(bomb) < 1000  # small on the wire, large once inflated
    with pytest.raises(ExtractError, match="unpacks to over"):
        extract_docx(bomb)


def test_zip_with_too_many_members_is_refused(monkeypatch):
    monkeypatch.setattr(extract, "MAX_ZIP_MEMBERS", 3)
    with pytest.raises(ExtractError, match="too many parts"):
        extract_bytes(_zip({f"f{i}": b"" for i in range(4)}), None, "x.pptx")


def test_non_zip_docx_is_an_extract_error():
    with pytest.raises(ExtractError, match="not a valid file"):
        extract_docx(b"not a zip")


def _fixture(name: str) -> bytes:
    """A file LibreOffice saved in the 97-2003 format, gzipped in the repo."""
    return gzip.decompress((Path(__file__).parent / "fixtures" / f"{name}.gz").read_bytes())


PPT_TEXT = ("--- Slide 1 ---\nDomain Name System\nDNS maps names to addresses\n"
            "Resolvers ask the root servers first\n\n"
            "--- Slide 2 ---\nRecord types\nA records hold IPv4 addresses\n"
            "MX records name mail exchangers – café")  # slide 3 has no text
DOC_TEXT = ("Indus Valley Civilization\nThe civilization flourished around 2500 BCE.\n"
            "Harappa and Mohenjo-daro were its largest cities – café.\n"
            "Period | Dates\nVedic | 1500 BCE")


@pytest.mark.parametrize("mime, name", [
    ("application/vnd.ms-powerpoint", "x"), (None, "Lecture%202.PPT"),
    ("application/octet-stream", "show.pps"),
])
def test_legacy_ppt_extracts_slide_text(mime, name):
    assert extract_bytes(_fixture("legacy.ppt"), mime, name) == PPT_TEXT


@pytest.mark.parametrize("mime, name", [("application/msword", "x"), (None, "essay.DOC")])
def test_legacy_doc_extracts_body_text_and_tables(mime, name):
    assert extract_bytes(_fixture("legacy.doc"), mime, name) == DOC_TEXT


def test_rtf_saved_as_doc_is_read():
    rtf = (rb"{\rtf1\ansi\ansicpg1252\uc1{\fonttbl{\f0 Arial;}}{\*\generator Word;}"
           rb"{\info{\title Hidden}}\pard Caf\'e9 \u8211?dash \u-10179?\u-8704? {\b bold}\par"
           b"\r\n"
           rb"{\field{\*\fldinst HYPERLINK \\l x}{\fldrslt link}}\par "
           rb"a\cell b\cell\row{\pict\bin4 }{}{}\}\{\\\par}")
    expected = "Café –dash \U0001f600 bold\nlink\na | b\n}{\\"
    assert extract_bytes(rtf, "application/msword", "notes.doc") == expected
    assert extract_bytes(rtf, "text/rtf", "x") == expected
    assert extract_bytes(rtf, None, "notes.rtf") == expected


def test_office_xml_under_a_legacy_name_goes_to_the_xml_reader(monkeypatch):
    monkeypatch.setattr(extract, "extract_pptx", lambda blob: "pptx")
    monkeypatch.setattr(extract, "extract_docx", lambda blob: "docx")
    blob = _zip({"[Content_Types].xml": b""})
    assert extract_bytes(blob, "application/vnd.ms-powerpoint", "a.ppt") == "pptx"
    assert extract_bytes(blob, "application/msword", "a.doc") == "docx"


@pytest.mark.parametrize("name, kind", [("a.ppt", "PPT"), ("a.doc", "DOC")])
def test_broken_legacy_files_are_extract_errors(name, kind):
    blobs = [b"not an office file", legacy.OLE_MAGIC + b"\0" * 600,
             _fixture("legacy.ppt")[:4096], _fixture("legacy.doc")[:4096]]
    for blob in blobs:
        with pytest.raises(ExtractError, match=f"{kind} can't be read"):
            extract_bytes(blob, None, name)


def test_a_doc_is_not_a_ppt_and_the_other_way_round():
    with pytest.raises(ExtractError, match="no 'PowerPoint Document' stream"):
        extract_bytes(_fixture("legacy.doc"), None, "a.ppt")
    with pytest.raises(ExtractError, match="no 'WordDocument' stream"):
        extract_bytes(_fixture("legacy.ppt"), None, "a.doc")


def test_legacy_files_under_a_modern_name_go_to_the_legacy_reader():
    pptx = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    assert extract_bytes(_fixture("legacy.ppt"), pptx, "x") == PPT_TEXT
    assert extract_bytes(_fixture("legacy.ppt"), None, "slides.pptx") == PPT_TEXT
    assert extract_bytes(_fixture("legacy.doc"), None, "essay.docx") == DOC_TEXT
    assert extract_bytes(rb"{\rtf1 hello\par}", None, "notes.docx") == "hello"


def _pptx(main_type: str | None = None) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Routing"
    buf = io.BytesIO()
    prs.save(buf)
    return _retyped(buf.getvalue(), b"presentationml.presentation.main+xml", main_type)


def _docx(main_type: str | None = None) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Switching")
    buf = io.BytesIO()
    doc.save(buf)
    return _retyped(buf.getvalue(), b"wordprocessingml.document.main+xml", main_type)


def _retyped(blob: bytes, old: bytes, main_type: str | None) -> bytes:
    """The package with its main part's content type swapped for `main_type`."""
    if main_type is None:
        return blob
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        members = {i.filename: z.read(i) for i in z.infolist()}
    full = b"application/vnd.openxmlformats-officedocument." + old
    assert full in members["[Content_Types].xml"]
    members["[Content_Types].xml"] = members["[Content_Types].xml"].replace(
        full, main_type.encode())
    return _zip(members)


OOXML = "application/vnd.openxmlformats-officedocument."


@pytest.mark.parametrize("main_type, mime, name", [
    (None, "application/octet-stream", "a.pptx"),
    (OOXML + "presentationml.slideshow.main+xml",
     OOXML + "presentationml.slideshow", "x"),
    (OOXML + "presentationml.slideshow.main+xml", "application/octet-stream", "a.PPSX"),
    (OOXML + "presentationml.template.main+xml", None, "a.potx"),
    ("application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml",
     "application/octet-stream", "a.pptm"),
    ("application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml",
     "application/vnd.ms-powerpoint.presentation.macroenabled.12", "x"),
])
def test_slide_shows_templates_and_macro_files_read_as_pptx(main_type, mime, name):
    assert extract_bytes(_pptx(main_type), mime, name) == "--- Slide 1 ---\nRouting"


@pytest.mark.parametrize("main_type, mime, name", [
    (None, "application/octet-stream", "a.docx"),
    ("application/vnd.ms-word.document.macroEnabled.main+xml",
     "application/octet-stream", "a.docm"),
    ("application/vnd.ms-word.document.macroEnabled.main+xml",
     "application/vnd.ms-word.document.macroenabled.12", "x"),
    (OOXML + "wordprocessingml.template.main+xml", None, "a.dotx"),
])
def test_templates_and_macro_files_read_as_docx(main_type, mime, name):
    assert extract_bytes(_docx(main_type), mime, name) == "Switching"


def test_opendocument_files_are_unsupported_not_fed_to_the_pptx_reader():
    odp = _zip({"content.xml": b"<x/>"})
    for mime, name in [("application/vnd.oasis.opendocument.presentation", "a.odp"),
                       ("application/vnd.oasis.opendocument.text", "a.odt")]:
        with pytest.raises(ExtractError, match="unsupported type"):
            extract_bytes(odp, mime, name)


def test_a_zip_that_is_no_office_package_is_an_extract_error():
    with pytest.raises(ExtractError, match="PPTX is not a PowerPoint file"):
        extract_bytes(_zip({"content.xml": b"<x/>"}), None, "a.pptx")
    with pytest.raises(ExtractError, match="DOCX is not a Word file") as e:
        extract_bytes(_pptx(), None, "a.docx")
    assert "BytesIO" not in str(e.value)


def _damaged(blob: bytes, member: str, method: int) -> bytes:
    """The package with `member`'s bytes garbled under an intact directory."""
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        members = {i.filename: z.read(i) for i in z.infolist()}
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", method) as z:
        for name, data in members.items():
            z.writestr(name, data)
        start = z.getinfo(member).header_offset + 30 + len(member.encode())
    raw = bytearray(buf.getvalue())
    for i in range(start + 2, start + 12):
        raw[i] ^= 0xFF
    return bytes(raw)


@pytest.mark.parametrize("method", [zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED])
@pytest.mark.parametrize("make, name, member", [
    (_pptx, "a.pptx", "[Content_Types].xml"),
    (_pptx, "a.pptx", "ppt/slides/slide1.xml"),
    (_docx, "a.docx", "[Content_Types].xml"),
    (_docx, "a.docx", "word/document.xml"),
])
def test_a_package_with_a_damaged_member_is_an_extract_error(make, name, member, method):
    kind = name[2:].upper()
    with pytest.raises(ExtractError, match=f"{kind} is damaged"):
        extract_bytes(_damaged(make(), member, method), None, name)


def _pdf(password: str | None = None) -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(200, 200)
    if password:
        writer.encrypt(password)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def test_unreadable_pdfs_are_extract_errors_in_plain_words():
    with pytest.raises(ExtractError, match="^PDF is password-protected$"):
        extract_bytes(_pdf("secret"), "application/pdf", "x.pdf")
    for blob in (b"%PDF-1.4 cut off", b"not a pdf", b""):
        with pytest.raises(ExtractError, match="PDF is damaged or incomplete"):
            extract_bytes(blob, "application/pdf", "x.pdf")
    with pytest.raises(ExtractError, match="no extractable text"):
        extract_bytes(_pdf(), "application/pdf", "x.pdf")


def test_a_pdf_is_known_by_its_first_bytes_whatever_it_is_called(monkeypatch):
    monkeypatch.setattr(extract, "extract_pdf", lambda blob: "pdf")
    assert extract_bytes(b"%PDF-1.7 ...", "application/msword", "notes.doc") == "pdf"
    assert extract_bytes(b"%PDF-1.7 ...", "application/octet-stream", "x") == "pdf"


def test_html_without_a_closing_head_tag_keeps_its_body():
    html = ("<!doctype html><html><head><meta charset=utf-8><title>T</title>"
            "<style>p{color:red}</style><body><p>hi</p><p>there</p></body></html>")
    assert html_to_text(html) == "hi\n\nthere"
    assert html_to_text("<head><title>T</title><p>no body tag</p>") == "no body tag"
    # a tag <head> may hold doesn't stop it ending at the first tag it may not
    assert html_to_text("<head><noscript>hidden<p>body") == "body"
    # non-whitespace after </title> ends an implicit <head>; title text stays skipped
    assert html_to_text("<head><title>T</title>bare") == "bare"
    assert html_to_text("<head><title>T</title>hello<p>world") == "hello\nworld"


def test_html_with_an_unclosed_skipped_element_keeps_the_rest_of_the_page():
    assert html_to_text("<p>a</p><noscript><img src=x><p>b</p><p>c</p>") == "a\n\nb\n\nc"
    # closing the outer element closes what was left open inside it
    assert html_to_text("<p>a</p><svg><template>x</svg><p>b</p>") == "a\n\nb"
    assert html_to_text("<p>a</p><script>var b = 1") == "a"


def test_legacy_files_with_no_text_are_extract_errors(monkeypatch):
    monkeypatch.setattr(legacy, "ppt_slides", lambda blob: ["", ""])
    monkeypatch.setattr(legacy, "doc_text", lambda blob: "")
    with pytest.raises(ExtractError, match="PPT has no extractable text"):
        extract_bytes(b"x", None, "a.ppt")
    with pytest.raises(ExtractError, match="DOC has no extractable text"):
        extract_bytes(b"x", None, "a.doc")


def _rec(rec_type: int, body: bytes, ver: int = 0, instance: int = 0) -> bytes:
    return struct.pack("<HHI", ver | instance << 4, rec_type, len(body)) + body


def _ppt_stream(saves: list[dict[int, bytes]], doc_id: int = 1) -> tuple[bytes, int]:
    """A "PowerPoint Document" stream: each save appends its records, a
    persist directory for them and a UserEditAtom pointing at the save before.
    Returns (stream, offset of the last UserEditAtom)."""
    data, last_edit = b"", 0
    for records in saves:
        entries = b""
        for persist_id, record in records.items():
            entries += struct.pack("<II", persist_id | 1 << 20, len(data))
            data += record
        directory = len(data)
        data += _rec(0x1772, entries)
        edit = len(data)
        data += _rec(0x0FF5, struct.pack("<IIIIIII", 0, 0, last_edit, directory, doc_id, 0, 0))
        last_edit = edit
    return data, last_edit


def _ppt_slides(monkeypatch, data: bytes, edit: int) -> list[str]:
    streams = {"PowerPoint Document": data,
               "Current User": _rec(0x0FF6, struct.pack("<III", 0x14, 0xE391C05F, edit))}
    monkeypatch.setattr(legacy, "_open", lambda blob: io.BytesIO())
    monkeypatch.setattr(legacy, "_stream", lambda ole, name: streams[name])
    return legacy.ppt_slides(b"")


def test_ppt_placeholder_text_is_read_from_the_outline_in_shape_order(monkeypatch):
    # what PowerPoint itself writes: title/body text in the document's slide
    # list, the slide's shapes pointing at it, text boxes carrying their own
    def header(kind):
        return _rec(0x0F9F, struct.pack("<I", kind))

    slide_list = _rec(0x0FF0, ver=0xF, body=(
        _rec(0x03F3, struct.pack("<IIiII", 2, 0, 2, 256, 0))
        + header(0) + _rec(0x0FA8, b"Title\x0bwrapped")
        + header(1) + _rec(0x0FA0, "Body – one\rtwo".encode("utf-16-le"))
        + header(1)  # an empty placeholder still takes an index
        + header(1) + _rec(0x0FA8, b"*")  # slide number
        + header(1) + _rec(0x0FA8, b"Unreferenced")
        + _rec(0x03F3, struct.pack("<IIiII", 3, 0, 0, 257, 0))
        + header(0) + _rec(0x0FA8, b"Second")))
    masters = _rec(0x0FF0, ver=0xF, instance=1, body=(
        _rec(0x03F3, struct.pack("<IIiII", 9, 0, 1, 1, 0))
        + header(0) + _rec(0x0FA8, b"Click to edit Master title style")))
    document = _rec(0x03E8, masters + slide_list, ver=0xF)

    def shape(*records):  # SpContainer > ClientTextbox, which isn't flagged a container
        return _rec(0xF004, _rec(0xF00D, b"".join(records)), ver=0xF)

    def ref(index):
        return _rec(0x0F9E, struct.pack("<i", index))

    slide = _rec(0x03EE, ver=0xF, body=_rec(0x040C, ver=0xF, body=(
        shape(ref(0)) + shape(header(4), _rec(0x0FA8, b"Text box")) + shape(ref(1))
        + shape(ref(2)) + shape(ref(3)) + shape(ref(99)))))
    data, edit = _ppt_stream([{1: document, 2: slide}])  # slide 3's record is missing
    assert _ppt_slides(monkeypatch, data, edit) == [
        "Title\nwrapped\nText box\nBody – one\ntwo\nUnreferenced", "Second"]


def test_ppt_saved_twice_reads_the_newest_copy_of_each_slide(monkeypatch):
    def document(*slide_ids):
        return _rec(0x03E8, ver=0xF, body=_rec(0x0FF0, ver=0xF, body=b"".join(
            _rec(0x03F3, struct.pack("<IIiII", i, 0, 0, 256, 0)) for i in slide_ids)))

    def slide(text):
        return _rec(0x03EE, _rec(0xF00D, _rec(0x0FA8, text)), ver=0xF)

    data, edit = _ppt_stream([
        {1: document(2), 2: slide(b"draft"), 3: slide(b"deleted later")},
        {1: document(2, 4), 2: slide(b"final"), 4: slide(b"added")},
    ])
    assert _ppt_slides(monkeypatch, data, edit) == ["final", "added"]


def test_ppt_password_protected_or_looping_is_refused(monkeypatch):
    data, edit = _ppt_stream([{1: _rec(0x03E8, b"", ver=0xF)}])
    streams = {"PowerPoint Document": data,
               "Current User": _rec(0x0FF6, struct.pack("<III", 0x14, 0xF3D1C4DF, edit))}
    monkeypatch.setattr(legacy, "_open", lambda blob: io.BytesIO())
    monkeypatch.setattr(legacy, "_stream", lambda ole, name: streams[name])
    with pytest.raises(ValueError, match="password-protected"):
        legacy.ppt_slides(b"")
    # a UserEditAtom whose "previous save" is itself must not loop forever
    loop = _rec(0x1772, b"") + _rec(0x0FF5, struct.pack("<IIIIIII", 0, 0, 8, 0, 1, 0, 0))
    with pytest.raises(ValueError, match="no document record"):
        _ppt_slides(monkeypatch, loop, 8)


def test_ppt_records_nested_past_the_limit_are_refused(monkeypatch):
    nested = _rec(0x0FA8, b"deep")
    for _ in range(legacy.MAX_DEPTH + 2):
        nested = _rec(0xF003, nested, ver=0xF)
    document = _rec(0x03E8, ver=0xF, body=_rec(0x0FF0, ver=0xF, body=_rec(
        0x03F3, struct.pack("<IIiII", 2, 0, 0, 256, 0))))
    data, edit = _ppt_stream([{1: document, 2: _rec(0x03EE, nested, ver=0xF)}])
    with pytest.raises(ValueError, match="nested too deep"):
        _ppt_slides(monkeypatch, data, edit)


def _doc(monkeypatch, pieces: list[tuple[str, bool]], body_chars: int, flags: int = 0x0200,
         ident: int = 0xA5EC) -> str:
    """A WordDocument + table stream holding `pieces` ((text, 8-bit?)) stored
    in reverse order, as a fast-saved file may."""
    word = bytearray(0x400)
    struct.pack_into("<H", word, 0, ident)
    struct.pack_into("<H", word, 0x0A, flags)
    struct.pack_into("<i", word, 0x4C, body_chars)
    cps, pcds, cp = [0], b"", 0
    stored = {}
    for text, compressed in reversed(pieces):
        stored[text] = len(word)
        word += text.encode("latin-1" if compressed else "utf-16-le")
    for text, compressed in pieces:
        cp += len(text)
        cps.append(cp)
        fc = stored[text] * 2 | 0x40000000 if compressed else stored[text]
        pcds += struct.pack("<HIH", 0, fc, 0)
    plc = struct.pack(f"<{len(cps)}I", *cps) + pcds
    table = b"\xff" * 7 + b"\x01" + struct.pack("<H", 3) + b"abc" \
        + b"\x02" + struct.pack("<I", len(plc)) + plc
    struct.pack_into("<II", word, 0x1A2, 7, len(table) - 7)
    streams = {"WordDocument": bytes(word), "1Table": table}
    monkeypatch.setattr(legacy, "_open", lambda blob: io.BytesIO())
    monkeypatch.setattr(legacy, "_stream", lambda ole, name: streams[name])
    return legacy.doc_text(b"")


def test_doc_pieces_are_joined_in_character_order_and_tidied(monkeypatch):
    body = [("Caf\xe9 \x93quoted\x94\rsoft\x1fhy\x1ephen\x0bline\x0c", True),
            ("हिन्दी \x13 HYPERLINK \"http://x\" \x14link\x15\r", False),
            ("a\x07b\x07\x07\x07c\x07\x07\x01\x08\r", True),
            ("\x13 PAGE \x13 nested \x14 1\x15\x15end\r", True)]
    footnotes = [("footnote text\r", True)]
    text = _doc(monkeypatch, body + footnotes, sum(len(t) for t, _ in body))
    assert text == ("Café “quoted”\nsofthy-phen\nline\nहिन्दी link\n"
                    "a | b\nc\n\nend")


def test_doc_unclosed_fields_dont_make_field_removal_quadratic():
    # 200k fields past their separator and never closed, then as much text:
    # hours if each kept character rescans the open fields
    n = 200_000
    text = "\x13\x14" * n + "x" * n + "\x13code\x14shown\x15\x15\x14y"
    start = time.monotonic()
    assert legacy._drop_field_codes(text) == "x" * n + "showny"
    assert time.monotonic() - start < 5


def test_doc_too_old_or_password_protected_is_refused(monkeypatch):
    with pytest.raises(ValueError, match="older than Word 97"):
        _doc(monkeypatch, [("x", True)], 1, ident=0xA5DC)
    with pytest.raises(ValueError, match="password-protected"):
        _doc(monkeypatch, [("x", True)], 1, flags=0x0300)


@pytest.mark.parametrize("url, video_id", [
    ("https://www.youtube.com/watch?v=dQw4w9WgXcQ", "dQw4w9WgXcQ"),
    ("https://m.youtube.com/watch?feature=share&v=dQw4w9WgXcQ", "dQw4w9WgXcQ"),
    ("https://youtu.be/dQw4w9WgXcQ?t=42", "dQw4w9WgXcQ"),
    ("https://www.youtube.com/shorts/dQw4w9WgXcQ", "dQw4w9WgXcQ"),
    ("https://www.youtube.com/embed/dQw4w9WgXcQ?start=3", "dQw4w9WgXcQ"),
    ("https://www.youtube-nocookie.com/embed/dQw4w9WgXcQ", "dQw4w9WgXcQ"),
    ("https://youtube.com/live/dQw4w9WgXcQ?si=abc", "dQw4w9WgXcQ"),
    (" https://youtu.be/dQw4w9WgXcQ ", "dQw4w9WgXcQ"),
    ("https://www.youtube.com/@lecturer/live", None),
    ("https://www.youtube.com/playlist?list=PL123", None),
    ("https://www.youtube.com/watch", None),
    ("https://notyoutu.be/dQw4w9WgXcQ", None),
    ("https://example.com/?next=youtu.be/dQw4w9WgXcQ", None),
    ("https://[bad", None),
    (None, None),
])
def test_youtube_video_id(url, video_id):
    assert youtube_video_id(url) == video_id


# -- Moodle download cap -------------------------------------------------------


class _Response:
    def __init__(self, body: bytes, length: str | None,
                 content_type: str = "application/pdf", url: str | None = None):
        self.body = io.BytesIO(body)
        self.reads = 0
        self.url = url or "https://m.example/webservice/pluginfile.php/1/a.pdf"
        headers = {"Content-Type": content_type}
        if length is not None:
            headers["Content-Length"] = length
        from email.message import Message

        self.headers = Message()
        for k, v in headers.items():
            self.headers[k] = v

    def geturl(self):
        return self.url

    def read(self, n=-1):
        self.reads += 1
        return self.body.read(n)

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


_PLUGINFILE = "https://m.example/webservice/pluginfile.php/1/a.pdf"


def _download(monkeypatch, response: _Response, fileurl: str = _PLUGINFILE):
    import app.moodle as moodle

    monkeypatch.setattr(moodle.urllib.request, "urlopen", lambda req, timeout: response)
    return MoodleClient("https://m.example", "tok").download(fileurl)


def test_download_refuses_large_content_length_without_reading(monkeypatch):
    import app.moodle as moodle

    monkeypatch.setattr(moodle, "MAX_DOWNLOAD_BYTES", 10)
    response = _Response(b"x" * 11, "11")
    with pytest.raises(ExtractError, match="over"):
        _download(monkeypatch, response)
    assert response.reads == 0


def test_download_stops_streaming_past_the_cap(monkeypatch):
    import app.moodle as moodle

    monkeypatch.setattr(moodle, "MAX_DOWNLOAD_BYTES", 10)
    with pytest.raises(ExtractError):  # permanent, not a retryable MoodleError
        _download(monkeypatch, _Response(b"x" * 11, None))


def test_download_within_the_cap(monkeypatch):
    blob, mime = _download(monkeypatch, _Response(b"%PDF-1.4 data", "13"))
    assert (blob, mime) == (b"%PDF-1.4 data", "application/pdf")


def test_download_network_errors_stay_retryable(monkeypatch):
    import app.moodle as moodle

    def boom(req, timeout):
        raise OSError("connection reset")

    monkeypatch.setattr(moodle.urllib.request, "urlopen", boom)
    with pytest.raises(MoodleError):
        MoodleClient("https://m.example", "tok").download(
            "https://m.example/pluginfile.php/1/a.pdf")


def test_download_keeps_an_html_file(monkeypatch):
    html = b"<!doctype html><p>lecture notes</p>"
    blob, mime = _download(monkeypatch, _Response(html, None, "text/html"))
    assert (blob, mime) == (html, "text/html")


def test_download_treats_a_redirect_off_pluginfile_as_a_login_page(monkeypatch):
    html = b"<!doctype html><html>log in</html>"
    response = _Response(html, None, "text/html", url="https://m.example/login/index.php")
    with pytest.raises(MoodleError, match="login page"):
        _download(monkeypatch, response)


def test_download_json_error_body_is_token_rejected_not_a_failed_pdf(monkeypatch):
    body = b'{"exception":"moodle_exception","errorcode":"invalidtoken","message":"Invalid token"}'
    with pytest.raises(TokenRejected) as exc:
        _download(monkeypatch, _Response(body, None, "application/json"))
    assert exc.value.errorcode == "invalidtoken"


def test_download_json_error_without_invalidtoken_is_retryable(monkeypatch):
    body = b'{"error":"No file","errorcode":"filenotfound"}'
    with pytest.raises(MoodleError) as exc:
        _download(monkeypatch, _Response(body, None, "application/json"))
    assert not isinstance(exc.value, TokenRejected)
    assert exc.value.errorcode == "filenotfound"


def test_download_keeps_a_json_file_even_when_it_looks_like_an_error(monkeypatch):
    body = b'{"error":"not moodle","errorcode":"invalidtoken"}'
    fileurl = "https://m.example/webservice/pluginfile.php/1/notes.json"
    response = _Response(body, None, "application/json", url=fileurl)
    blob, mime = _download(monkeypatch, response, fileurl)
    assert (blob, mime) == (body, "application/json")
