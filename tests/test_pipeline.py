"""Phase 2 tests: extraction dispatch + chunking (FakeLLM, no network)."""

import io

import pytest
from sqlmodel import Session, SQLModel, select, update

from app.chunk import chunk_resource, locate, presplit
from app.extract import (
    ExtractError,
    SkipResource,
    extract_bytes,
    extract_docx,
    extract_pptx,
    extract_resource_text,
    extract_transcript,
)
from app.models import Chunk, Course, QuizItem, Resource, Topic
from app.moodle import MoodleError
from app.pipeline import (
    pending_resource_ids,
    quiz_chunk_ids,
    run_chunking,
    run_extraction,
    single_run,
)
from app.sync import ContentChanged, commit_if_current
from tests.dbutil import TEST_DATABASE_URL, make_engine


class FakeLLM:
    def complete_json(self, system, user, **kw):
        # verbatim halves of the section -> offsets must resolve
        text = user.split("\n\n", 1)[1]
        half = len(text) // 2
        return {"chunks": [
            {"title": "C0", "content": text[:half]},
            {"title": "C1", "content": text[half:]},
        ]}


@pytest.fixture()
def session():
    engine = make_engine("sqlite:///:memory:")
    SQLModel.metadata.create_all(engine)
    with Session(engine) as s:
        yield s


def _resource(session, **kw):
    course = session.exec(
        select(Course).where(Course.source_id == "c1")).first()
    if course is None:
        course = Course(source="moodle", source_id="c1", name="C")
        session.add(course)
        session.commit()
    topic = session.exec(
        select(Topic).where(Topic.course_id == course.id)).first()
    if topic is None:
        topic = Topic(course_id=course.id, source_id="t1", title="T")
        session.add(topic)
        session.commit()
    kw.setdefault("topic_id", topic.id)
    kw.setdefault("source", "moodle")
    kw.setdefault("source_id", "r1")
    kw.setdefault("type", "file")
    kw.setdefault("title", "R")
    kw.setdefault("status", "pending")
    r = Resource(**kw)
    session.add(r)
    session.commit()
    session.refresh(r)
    return r


def _docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Hello docx world")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx_bytes() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 10, 10)
    box.text_frame.text = "Hello pptx world"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_docx_roundtrip():
    assert "Hello docx world" in extract_docx(_docx_bytes())


def test_pptx_roundtrip():
    text = extract_pptx(_pptx_bytes())
    assert "Slide 1" in text and "Hello pptx world" in text


def test_dispatch_and_unsupported():
    assert "Hello" in extract_bytes(_docx_bytes(), None, "notes.docx")
    assert "plain" in extract_bytes(b"plain text", "text/plain", "n.txt")
    with pytest.raises(ExtractError):
        extract_bytes(b"\x00\x01", "application/x-unknown", "n.bin")


def test_video_without_url_skipped(session):
    r = _resource(session, type="video", raw_url=None)
    with pytest.raises(SkipResource):
        extract_resource_text(r)


def test_link_skipped(session):
    r = _resource(session, type="link", raw_url="https://example.com")
    with pytest.raises(SkipResource):
        extract_resource_text(r)


class _Snippet:
    def __init__(self, text):
        self.text = text


class _Transcript:
    def __init__(self, lang, generated=False, translatable=True):
        self.language_code, self.is_generated = lang, generated
        self.is_translatable = translatable

    def translate(self, lang):
        return _Transcript(f"{self.language_code}->{lang}", self.is_generated)

    def fetch(self):
        return [_Snippet(f"text in {self.language_code}")]


class _UnfetchableTranslation(_Transcript):
    def translate(self, lang):
        import youtube_transcript_api as yta

        class Broken(_Transcript):
            def fetch(self):
                raise yta.YouTubeTranscriptApiException("translation request failed")

        return Broken(lang)


def _fake_youtube(monkeypatch, available):
    """Fake API over `available` transcripts; returns the list() call log."""
    import youtube_transcript_api as yta

    listed = []

    class FakeList:
        def __iter__(self):
            return iter(available)

        def find_transcript(self, languages):
            for lang in languages:  # manual before generated, like the real one
                for generated in (False, True):
                    for t in available:
                        if t.language_code == lang and t.is_generated == generated:
                            return t
            raise yta.NoTranscriptFound("abc", languages, None)

    class FakeApi:
        def list(self, video_id):
            listed.append(video_id)
            return FakeList()

        def fetch(self, *a, **k):
            raise AssertionError("fetch() lists again; use the one list")

    monkeypatch.setattr(yta, "YouTubeTranscriptApi", FakeApi)
    return listed


def test_transcript_prefers_english(monkeypatch):
    listed = _fake_youtube(monkeypatch, [_Transcript("fr"), _Transcript("en", generated=True),
                                         _Transcript("en-GB")])
    assert extract_transcript("https://youtu.be/abc") == "text in en"
    assert listed == ["abc"]


def test_transcript_language_order_beats_manual_vs_generated(monkeypatch):
    _fake_youtube(monkeypatch, [_Transcript("en", generated=True),
                                _Transcript("en-US")])
    # "en" is tried first, generated "en" beats manual "en-US" (language order wins)
    assert extract_transcript("https://youtu.be/abc") == "text in en"


def test_transcript_falls_back_to_translated_manual_captions(monkeypatch):
    listed = _fake_youtube(monkeypatch, [_Transcript("sw", generated=True), _Transcript("fr")])
    assert extract_transcript("https://www.youtube.com/watch?v=abc") == "text in fr->en"
    assert listed == ["abc"]  # one caption-list request, not two


def test_transcript_untranslatable_uses_original_language(monkeypatch):
    _fake_youtube(monkeypatch, [_Transcript("sw", translatable=False)])
    assert extract_transcript("https://youtu.be/abc") == "text in sw"


def test_transcript_failed_translation_fetch_falls_back_to_original(monkeypatch):
    _fake_youtube(monkeypatch, [_UnfetchableTranslation("fr")])
    assert extract_transcript("https://youtu.be/abc") == "text in fr"


def test_transcript_none_available_skips(monkeypatch):
    _fake_youtube(monkeypatch, [])
    with pytest.raises(SkipResource):
        extract_transcript("https://youtu.be/abc")


def test_presplit_and_locate():
    text = ("para one\n\n" * 5000)
    assert len(presplit(text)) > 1
    assert locate("b", "abc") == (1, 2)
    assert locate("zzz", "abc") == (None, None)


def test_locate_fuzzy_whitespace():
    full = "Identifying  concepts\nrelated  to  networks"
    assert locate("Identifying concepts related to networks", full) == (0, 43)


def test_chunk_short_text_no_llm(session):
    r = _resource(session, status="extracted", extracted_text="tiny but real")
    assert chunk_resource(session, r, None) == 1
    chunk = select(Chunk)
    assert len(session.exec(chunk).all()) == 1


def test_chunk_regenerates_on_change(session):
    r = _resource(session, status="extracted", extracted_text="x" * 500)
    assert chunk_resource(session, r, FakeLLM()) == 2
    r.status = "pending"  # sync saw new content
    session.add(r)
    session.commit()
    assert chunk_resource(session, r, FakeLLM()) == 2
    assert len(session.exec(select(Chunk)).all()) == 2  # replaced, not doubled


def test_chunk_cached(session):
    r = _resource(session, status="extracted", extracted_text="x" * 500)
    assert chunk_resource(session, r, FakeLLM()) == 2
    assert chunk_resource(session, r, FakeLLM()) == 0


def test_run_extraction_counts(session):
    _resource(session, source_id="ok", type="page_text", text="page content here",
              extracted_text="page content here")
    _resource(session, source_id="skip", type="link", raw_url="https://example.com")
    counts = run_extraction(session, downloader=None).counts
    assert counts == {"extracted": 1, "skipped": 1, "failed": 0}


def test_run_extraction_only_touches_its_source(session):
    # a Moodle downloader appends the Moodle token: it must never see Drive URLs
    _resource(session, source_id="m", raw_url="https://m.example/notes.txt")
    _resource(session, source_id="g", source="classroom",
              raw_url="https://drive.google.com/file/d/x/view")
    seen = []

    def downloader(url):
        seen.append(url)
        return b"notes", "text/plain"

    counts = run_extraction(session, downloader, source="moodle").counts
    assert counts["extracted"] == 1
    assert seen == ["https://m.example/notes.txt"]
    by_id = {r.source_id: r for r in session.exec(select(Resource)).all()}
    assert by_id["g"].status == "pending"


def test_run_extraction_fails_foreign_urls(session):
    from app.moodle import MoodleClient

    _resource(session, source_id="g", raw_url="https://drive.google.com/file/d/x/view")
    client = MoodleClient("https://m.example", "TOKEN")
    counts = run_extraction(session, client.download).counts
    assert counts["failed"] == 1
    r = session.exec(select(Resource)).one()
    assert r.status == "failed" and "refusing" in r.error and "TOKEN" not in r.error


def test_run_extraction_survives_download_errors(session):
    _resource(session, source_id="down", raw_url="https://m.example/down.txt")
    _resource(session, source_id="bad", raw_url="https://m.example/bad.bin")
    _resource(session, source_id="ok", raw_url="https://m.example/ok.txt")

    def downloader(url):
        if "down" in url:
            raise MoodleError("download failed: timed out")
        if "bad" in url:
            raise ValueError("parser blew up")  # stands in for a library crash
        return b"plain text notes", "text/plain"

    counts = run_extraction(session, downloader).counts
    assert counts == {"extracted": 1, "skipped": 0, "failed": 1, "download_errors": 1}
    by_id = {r.source_id: r for r in session.exec(select(Resource)).all()}
    assert by_id["down"].status == "pending"  # retried next run
    assert "timed out" in by_id["down"].error
    assert by_id["bad"].status == "failed" and "ValueError" in by_id["bad"].error
    assert by_id["ok"].status == "extracted"


def test_run_chunking_counts(session):
    _resource(session, source_id="a", status="extracted", extracted_text="short one")
    counts = run_chunking(session, None).counts
    assert counts["chunks"] == 1 and counts["resources"] == 1
    assert run_chunking(session, None).counts["cached"] == 1


def test_chunk_without_llm_leaves_long_text_alone(session):
    r = _resource(session, status="extracted", extracted_text="x" * 500)
    assert chunk_resource(session, r, FakeLLM()) == 2
    r.status = "pending"  # content changed; stale chunks await an LLM run
    session.add(r)
    session.commit()
    assert chunk_resource(session, r, None) == 0
    session.refresh(r)
    assert r.status == "pending" and r.error is None
    assert len(session.exec(select(Chunk)).all()) == 2  # not deleted


def test_run_chunking_without_llm_skips_not_fails(session):
    _resource(session, source_id="long", status="extracted", extracted_text="x" * 500)
    result = run_chunking(session, None)
    assert result.counts["needs_llm"] == 1 and result.counts["chunks"] == 0
    r = session.exec(select(Resource)).one()
    assert r.status == "extracted" and r.error is None


def test_stage_ids_are_scoped_to_course(session):
    mine = _resource(session, source_id="m", status="extracted", extracted_text="t")
    other_course = Course(source="moodle", source_id="c2", name="Other")
    session.add(other_course)
    session.commit()
    other_topic = Topic(course_id=other_course.id, source_id="t2", title="T2")
    session.add(other_topic)
    session.commit()
    theirs = _resource(session, source_id="o", topic_id=other_topic.id,
                       status="extracted", extracted_text="t")
    for res in (mine, theirs):
        session.add(Chunk(resource_id=res.id, title="c", content="t", order=0))
    session.commit()
    mine_chunk = session.exec(select(Chunk).where(Chunk.resource_id == mine.id)).one()
    my_course_id = session.get(Topic, mine.topic_id).course_id
    assert quiz_chunk_ids(session, my_course_id) == [mine_chunk.id]
    assert len(quiz_chunk_ids(session)) == 2


def _classroom_file(session, owner_email):
    from app.models import User

    user = User(email=owner_email)
    session.add(user)
    session.commit()
    course = Course(user_id=user.id, source="classroom", source_id="gc1", name="GC")
    session.add(course)
    session.commit()
    topic = Topic(course_id=course.id, source_id="t", title="T")
    session.add(topic)
    session.commit()
    return _resource(session, topic_id=topic.id, source="classroom", source_id="m:0",
                     raw_url="https://drive.google.com/file/d/1AbCdEfGhIjKlMnOp/view")


def _fake_drive(monkeypatch, download):
    import app.auth
    import app.drive

    monkeypatch.setattr(app.auth, "classroom_token_for",
                        lambda user: "rt" if user.email == "has@x.edu" else None)
    monkeypatch.setattr(app.drive, "build_service", lambda *a: None)
    monkeypatch.setattr(app.drive.DriveClient, "download", lambda self, url: download(url))


def test_classroom_drive_files_are_extracted_as_owner(session, monkeypatch):
    from app.pipeline import DOWNLOADERS_FOR

    _fake_drive(monkeypatch, lambda url: (b"lecture notes", "text/plain"))
    _classroom_file(session, "has@x.edu")
    counts = run_extraction(session, None, source="classroom",
                            downloader_for=DOWNLOADERS_FOR["classroom"](session)).counts
    assert counts["extracted"] == 1
    r = session.exec(select(Resource)).one()
    assert (r.status, r.extracted_text) == ("extracted", "lecture notes")


def test_classroom_owner_without_token_or_drive_grant_stays_pending(session, monkeypatch):
    from app.drive import DriveError
    from app.pipeline import DOWNLOADERS_FOR

    def no_grant(url):
        raise DriveError("token lacks Drive access; the owner must sign in again")

    _fake_drive(monkeypatch, no_grant)
    _classroom_file(session, "none@x.edu")
    counts = run_extraction(session, None, source="classroom",
                            downloader_for=DOWNLOADERS_FOR["classroom"](session)).counts
    assert counts["no_token"] == 1
    # the owner signs in (now has a token) but it predates the Drive grant
    from app.models import User

    owner = session.exec(select(User)).one()
    owner.email = "has@x.edu"
    session.add(owner)
    session.commit()
    counts = run_extraction(session, None, source="classroom",
                            downloader_for=DOWNLOADERS_FOR["classroom"](session)).counts
    assert counts["download_errors"] == 1
    r = session.exec(select(Resource)).one()
    assert r.status == "pending" and "sign in again" in r.error


# -- issue #27: presplit fallbacks, truncation, build-then-swap, backoff ------


def test_presplit_falls_back_to_lines_then_spaces_then_slices():
    lines = "a transcript line\n" * 2000  # no blank lines at all
    parts = presplit(lines, 1000)
    assert len(parts) > 1 and all(len(p) <= 1000 for p in parts)
    words = "word " * 3000
    assert all(len(p) <= 1000 for p in presplit(words, 1000))
    blob = "x" * 2500
    assert presplit(blob, 1000) == ["x" * 1000, "x" * 1000, "x" * 500]
    assert presplit("   \n\n  ", 1000) == []


class TruncatingLLM(FakeLLM):
    """Truncates any section longer than `limit`, like a capped max_tokens."""

    def __init__(self, limit):
        self.limit, self.sizes = limit, []

    def complete_json(self, system, user, **kw):
        from app.llm import TruncatedError

        size = len(user.split("\n\n", 1)[1])
        self.sizes.append(size)
        if size > self.limit:
            raise TruncatedError("finish_reason=length")
        return super().complete_json(system, user, **kw)


def test_truncated_section_is_halved_and_retried():
    from app.chunk import chunk_sections

    llm = TruncatingLLM(1500)
    out = chunk_sections(["para\n\n" * 500], llm)  # 3000 chars
    assert out and llm.sizes[0] == 3000
    assert all(s <= 1500 for s in llm.sizes[1:])


def test_truncation_on_a_small_section_raises():
    from app.chunk import chunk_sections
    from app.llm import TruncatedError

    with pytest.raises(TruncatedError):
        chunk_sections(["x" * 500], TruncatingLLM(100))


class MessyLLM:
    def __init__(self, chunks):
        self.chunks = chunks

    def complete_json(self, system, user, **kw):
        assert kw.get("required_key") == "chunks"
        return {"chunks": self.chunks}


def test_malformed_chunks_are_dropped_not_crashed_on():
    from app.chunk import chunk_sections

    good = {"title": "T", "content": "real text"}
    assert chunk_sections(["s"], MessyLLM("oops")) == []
    out = chunk_sections(["s"], MessyLLM([
        "a string", None, {"title": "no content"}, {"content": None},
        {"content": ["list"]}, {"title": None, "content": "untitled ok"}, good,
    ]))
    assert out == [{"title": "Untitled", "content": "untitled ok"}, good]


class FailingLLM:
    calls = 0

    def complete_json(self, system, user, **kw):
        type(self).calls += 1
        raise RuntimeError("provider exploded")


def test_failed_rechunk_keeps_the_old_chunks(session):
    from app.chunk import ChunkingError

    r = _resource(session, status="extracted", extracted_text="x" * 500)
    chunk_resource(session, r, FakeLLM())
    r.status = "pending"
    session.add(r)
    session.commit()
    with pytest.raises(RuntimeError):
        chunk_resource(session, r, FailingLLM())
    session.rollback()
    assert len(session.exec(select(Chunk)).all()) == 2
    with pytest.raises(ChunkingError):
        chunk_resource(session, r, MessyLLM([]))
    session.rollback()
    assert len(session.exec(select(Chunk)).all()) == 2


def test_blank_text_is_skipped_without_a_chunk(session):
    r = _resource(session, status="extracted", extracted_text="  \n ")
    assert chunk_resource(session, r, None) == 0
    session.refresh(r)
    assert r.status == "skipped" and not session.exec(select(Chunk)).all()


def test_extraction_of_empty_text_is_skipped(session):
    _resource(session, source_id="e", raw_url="https://m.example/empty.txt")
    counts = run_extraction(session, lambda url: (b"   ", "text/plain")).counts
    assert counts["skipped"] == 1 and counts["extracted"] == 0
    assert session.exec(select(Resource)).one().status == "skipped"


def _age(session, r):
    """Pretend the backoff has elapsed."""
    from datetime import datetime, timedelta, timezone

    r.retry_after = datetime.now(timezone.utc) - timedelta(seconds=1)
    session.add(r)
    session.commit()


def test_chunking_failure_backs_off_then_gives_up(session):
    from app.pipeline import MAX_CHUNK_ATTEMPTS, chunkable_resource_ids

    r = _resource(session, status="extracted", extracted_text="x" * 500)
    FailingLLM.calls = 0
    for attempt in range(1, MAX_CHUNK_ATTEMPTS + 1):
        assert run_chunking(session, FailingLLM()).counts["errors"] == 1
        session.refresh(r)
        assert r.attempts == attempt and r.retry_after is not None
        # deferred: the very next run doesn't bill it again
        assert run_chunking(session, FailingLLM()).counts["errors"] == 0
        _age(session, r)
    assert FailingLLM.calls == MAX_CHUNK_ATTEMPTS
    assert r.status == "failed" and "provider exploded" in r.error
    assert chunkable_resource_ids(session) == []


def test_backoff_doubles_and_is_capped():
    from datetime import datetime, timedelta, timezone

    from app.pipeline import _defer

    r = Resource(topic_id=None, source="moodle", source_id="x", type="file", title="R")
    gaps = []
    for _ in range(7):
        _defer(r, "e")
        gaps.append(round((r.retry_after - datetime.now(timezone.utc)) / timedelta(hours=1)))
    assert gaps == [1, 2, 4, 8, 16, 24, 24]


def test_success_after_a_failure_resets_the_backoff(session):
    r = _resource(session, status="extracted", extracted_text="x" * 500)
    run_chunking(session, FailingLLM())
    _age(session, r)
    assert run_chunking(session, FakeLLM()).counts["chunks"] == 2
    session.refresh(r)
    assert r.attempts == 0 and r.retry_after is None and r.status == "extracted"


def test_download_errors_back_off_but_stay_pending(session):
    from app.pipeline import pending_resource_ids

    r = _resource(session, source_id="down", raw_url="https://m.example/down.txt")

    def down(url):
        raise MoodleError("timed out")

    assert run_extraction(session, down).counts["download_errors"] == 1
    session.refresh(r)
    assert r.status == "pending" and r.attempts == 1
    assert pending_resource_ids(session) == []  # not retried until due
    _age(session, r)
    assert run_extraction(session, lambda url: (b"notes", "text/plain")).counts["extracted"] == 1
    session.refresh(r)
    assert r.attempts == 0 and r.retry_after is None


def _sync_replaces(session, rid):
    """What sync does on a content_hash change, committed mid-call the way a
    worker sync job commits while the pipeline waits on a download or LLM."""
    from app.sync import _purge_derived, lock_resource

    lock_resource(session, rid)
    _purge_derived(session, rid)
    session.exec(update(Resource).where(Resource.id == rid).values(
        content_hash="new", extracted_text=None, status="pending"))
    session.commit()


class RacingLLM(FakeLLM):
    def __init__(self, session, rid):
        self.session, self.rid = session, rid

    def complete_json(self, system, user, **kw):
        _sync_replaces(self.session, self.rid)
        return super().complete_json(system, user, **kw)


def _assert_reset(session, rid):
    session.expire_all()
    r = session.get(Resource, rid)
    assert (r.content_hash, r.status, r.extracted_text) == ("new", "pending", None)
    assert session.exec(select(Chunk)).all() == []
    assert rid in pending_resource_ids(session)  # extracted again next run


def test_chunks_of_text_replaced_mid_call_are_dropped(session):
    r = _resource(session, status="extracted", extracted_text="x" * 500,
                  content_hash="old")
    counts = run_chunking(session, RacingLLM(session, r.id)).counts
    assert counts["changed"] == 1 and counts["chunks"] == 0
    _assert_reset(session, r.id)


def test_rechunk_replaced_mid_call_keeps_neither_old_nor_new_chunks(session):
    r = _resource(session, status="extracted", extracted_text="x" * 500,
                  content_hash="old")
    assert chunk_resource(session, r, FakeLLM()) == 2
    r.status = "pending"
    session.add(r)
    session.commit()
    with pytest.raises(ContentChanged):
        chunk_resource(session, r, RacingLLM(session, r.id))
    _assert_reset(session, r.id)


def test_text_downloaded_before_a_content_change_is_dropped(session):
    r = _resource(session, raw_url="https://m.example/notes.txt", content_hash="old")
    rid = r.id

    def downloader(url):
        _sync_replaces(session, rid)
        return b"old notes", "text/plain"

    counts = run_extraction(session, downloader).counts
    assert counts["changed"] == 1 and counts["extracted"] == 0
    _assert_reset(session, rid)


def test_failure_on_replaced_content_does_not_back_off_the_new_one(session):
    r = _resource(session, status="extracted", extracted_text="x" * 500,
                  content_hash="old")

    class RacingFailingLLM:
        def complete_json(self, *a, **kw):
            _sync_replaces(session, r.id)
            raise RuntimeError("provider 500")

    counts = run_chunking(session, RacingFailingLLM()).counts
    assert counts["changed"] == 1 and "errors" not in counts
    session.expire_all()
    fresh = session.get(Resource, r.id)
    assert fresh.attempts == 0 and fresh.retry_after is None


def test_commit_if_current_rolls_back_on_a_new_hash(session):
    r = _resource(session, content_hash="old")
    r.title = "edited"
    session.add(r)
    assert commit_if_current(session, r.id, "old")
    r.title = "lost"
    session.add(r)
    assert not commit_if_current(session, r.id, "other")
    session.expire_all()
    assert session.get(Resource, r.id).title == "edited"


def test_rechunk_drops_quiz_items_of_the_old_chunks(session):
    r = _resource(session, status="extracted", extracted_text="x" * 500)
    chunk_resource(session, r, FakeLLM())
    chunk = session.exec(select(Chunk)).first()
    session.add(QuizItem(chunk_id=chunk.id, question="Q?", question_type="mcq",
                         options=["a", "b"], correct_answer="0",
                         generation_key="g-rechunk"))
    session.commit()
    r.status = "pending"
    session.add(r)
    session.commit()
    assert chunk_resource(session, r, FakeLLM()) == 2
    assert session.exec(select(QuizItem)).all() == []  # no orphans on SQLite


def test_single_run_on_sqlite_always_runs(session):
    with single_run("moodle", session.get_bind()) as got:
        assert got


@pytest.mark.skipif(not TEST_DATABASE_URL, reason="advisory locks need Postgres")
def test_single_run_excludes_a_second_run_of_the_same_source():
    a, b = make_engine(), make_engine()
    with single_run("moodle", a) as first:
        assert first
        with single_run("moodle", b) as second:
            assert not second
        with single_run("classroom", b) as other:
            assert other  # sources touch disjoint resources
    with single_run("moodle", b) as again:
        assert again  # released on exit


@pytest.mark.parametrize("flag, stages, needs_llm", [
    (None, ("extraction", "chunking", "quiz"), True),
    ("--extract-only", ("extraction",), False),
    ("--chunk-only", ("chunking",), True),
    ("--quiz-only", ("quiz",), True),
])
def test_main_runs_the_stages_its_flags_pick(session, monkeypatch, flag, stages, needs_llm):
    import app.pipeline as pipeline

    ran: list = []
    monkeypatch.setattr(pipeline, "engine", session.get_bind())
    monkeypatch.setattr(pipeline, "llm_clients", lambda: ("chunk", "quiz"))
    monkeypatch.setattr(pipeline, "run_course",
                        lambda s, source, cid, c, q, pace, stages: ran.append((c, q, stages)))
    monkeypatch.setattr("sys.argv", ["pipeline", "--source", "moodle", *([flag] if flag else [])])
    pipeline.main()
    llms = ("chunk", "quiz") if needs_llm else (None, None)
    assert [(c, q, tuple(st)) for c, q, st in ran] == [(*llms, stages)]


@pytest.mark.parametrize("job_running", [True, False])
def test_main_waits_only_behind_a_pipeline_job(session, monkeypatch, capsys, job_running):
    from contextlib import contextmanager

    import app.pipeline as pipeline
    from app.jobs import enqueue

    if job_running:
        job = enqueue(session, "pipeline", {"source": "moodle", "user_email": "s@x.edu"})
        job.status = "running"
        session.add(job)
        session.commit()
    calls, ran = [], []

    @contextmanager
    def single_run(source, bind=None, wait=False):
        calls.append(wait)
        yield wait  # the try fails; a blocking wait gets the lock

    monkeypatch.setattr(pipeline, "engine", session.get_bind())
    monkeypatch.setattr(pipeline, "single_run", single_run)
    monkeypatch.setattr(pipeline, "run_course", lambda *a, **kw: ran.append(1))
    monkeypatch.setattr("sys.argv", ["pipeline", "--source", "moodle", "--extract-only"])
    pipeline.main()
    if job_running:
        assert calls == [False, True] and ran == [1]
        assert "waiting for it" in capsys.readouterr().out
    else:
        assert calls == [False] and ran == []
        assert "in progress; exiting" in capsys.readouterr().out


@pytest.mark.skipif(not TEST_DATABASE_URL, reason="advisory locks need Postgres")
def test_single_run_wait_blocks_until_the_holder_releases():
    import threading

    a, b = make_engine(), make_engine()
    held, release, got = threading.Event(), threading.Event(), []

    def holder():
        with single_run("moodle", a) as first:
            assert first
            held.set()
            release.wait(5)

    t = threading.Thread(target=holder)
    t.start()
    held.wait(5)

    def waiter():
        with single_run("moodle", b, wait=True) as second:
            got.append(second)

    w = threading.Thread(target=waiter)
    w.start()
    w.join(0.5)
    assert got == []  # still blocked
    release.set()
    w.join(5)
    t.join(5)
    assert got == [True]


def test_quiz_candidates_exclude_attempted_chunks(session):
    from app.models import QuizAttempt

    r = _resource(session, source_id="q", status="extracted", extracted_text="t")
    chunks = [Chunk(resource_id=r.id, title=f"c{i}", content="t", order=i) for i in range(3)]
    session.add_all(chunks)
    session.commit()
    session.add_all([QuizAttempt(chunk_id=chunks[0].id, attempt=1),
                     QuizAttempt(chunk_id=chunks[1].id, attempt=2)])
    session.commit()
    assert quiz_chunk_ids(session, attempt=1) == [chunks[1].id, chunks[2].id]
    assert quiz_chunk_ids(session, attempt=2) == [chunks[0].id, chunks[2].id]
    assert len(quiz_chunk_ids(session)) == 3


class FlakyQuizLLM:
    """Fails the first `fail` calls (an unparseable reply), then succeeds."""

    def __init__(self, fail):
        self.fail, self.calls = fail, 0

    def complete_json(self, system, user, **kw):
        self.calls += 1
        if self.calls <= self.fail:
            from app.llm import LLMError

            raise LLMError("reply was not valid JSON")
        return {"items": [{"question": "Q?", "question_type": "short_answer",
                           "correct_answer": "A", "grading_criteria": "A",
                           "difficulty": "recall"}]}


def _quiz_chunk(session):
    r = _resource(session, source_id="qf", status="extracted", extracted_text="t")
    chunk = Chunk(resource_id=r.id, title="c", content="t", order=0)
    session.add(chunk)
    session.commit()
    return chunk


def _expire_backoff(session, chunk):
    from datetime import datetime, timezone

    from app.models import QuizFailure

    session.exec(update(QuizFailure).where(QuizFailure.chunk_id == chunk.id)
                 .values(retry_after=datetime(2000, 1, 1, tzinfo=timezone.utc)))
    session.commit()


def test_failed_quiz_generation_backs_off_then_gives_up(session):
    from app.models import QuizFailure
    from app.pipeline import MAX_QUIZ_FAILURES, run_quiz

    chunk = _quiz_chunk(session)
    llm = FlakyQuizLLM(fail=99)
    assert run_quiz(session, llm).counts["errors"] == 1
    f = session.get(QuizFailure, (chunk.id, 1))
    assert f.failures == 1 and f.retry_after is not None and f.error.startswith("LLMError")
    # backed off: the next run doesn't call the LLM again
    assert quiz_chunk_ids(session, attempt=1) == []
    run_quiz(session, llm)
    assert llm.calls == 1
    for _ in range(MAX_QUIZ_FAILURES - 1):
        _expire_backoff(session, chunk)
        res = run_quiz(session, llm)
    assert llm.calls == MAX_QUIZ_FAILURES and res.counts["given_up"] == 1
    _expire_backoff(session, chunk)  # given up even once the backoff is over
    assert quiz_chunk_ids(session, attempt=1) == []
    # another attempt number is tracked separately
    assert quiz_chunk_ids(session, attempt=2) == [chunk.id]


def test_quiz_success_after_failure_clears_the_backoff(session):
    from app.models import QuizAttempt, QuizFailure
    from app.pipeline import run_quiz

    chunk = _quiz_chunk(session)
    llm = FlakyQuizLLM(fail=1)
    run_quiz(session, llm)
    _expire_backoff(session, chunk)
    assert run_quiz(session, llm).counts["items"] == 1
    session.expire_all()
    assert session.get(QuizFailure, (chunk.id, 1)) is None
    assert session.get(QuizAttempt, (chunk.id, 1)) is not None


def test_quiz_paces_after_failed_calls(session, monkeypatch):
    from app.pipeline import run_quiz

    for i in range(2):
        r = _resource(session, source_id=f"p{i}", status="extracted", extracted_text="t")
        session.add(Chunk(resource_id=r.id, title="c", content=f"t{i}", order=0))
    session.commit()
    slept = []
    monkeypatch.setattr("app.pipeline.time.sleep", slept.append)
    assert run_quiz(session, FlakyQuizLLM(fail=99), pace=2.5).counts["errors"] == 2
    assert slept == [2.5, 2.5]


def test_resync_purges_quiz_failures(session):
    from app.models import QuizFailure
    from app.pipeline import run_quiz
    from app.sync import _purge_derived

    chunk = _quiz_chunk(session)
    run_quiz(session, FlakyQuizLLM(fail=99))
    _purge_derived(session, chunk.resource_id)
    session.commit()
    assert session.exec(select(QuizFailure)).all() == []
